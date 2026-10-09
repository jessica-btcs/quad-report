#!/usr/bin/env python3
"""Layout check for a built Quad slide: is every word inside its box and on the page?

    python3 layout_check.py <slide.pptx>

build_quad.py runs this after every build. It measures each text box with
Calibri's own character widths (embedded below, so the answer is the same on
any machine, with or without Calibri installed), wraps each paragraph the way
PowerPoint does at that width, and compares the height the text needs with
the height the box has. It also checks every box sits inside the slide.

Prints JSON: {"ok": bool, "problems": [...]}. Exit 0 clean, 3 if anything
overflows. Each problem names the quadrant, the start of the text, and how
much is over, so the fix is to shorten that text and build again.
"""
import json, sys

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    sys.exit("This needs python-pptx.  Install it with:  pip install python-pptx")

# Advance widths as a fraction of the font size, from calibri.ttf,
# calibrib.ttf and calibrii.ttf (2048 units per em).
REGULAR = {
    " ": 0.2261, "!": 0.3257, "\"": 0.4009, "#": 0.498, "$": 0.5068, "%": 0.7148,
    "&": 0.6821, "'": 0.2207, "(": 0.3032, ")": 0.3032, "*": 0.498, "+": 0.498,
    ",": 0.2495, "-": 0.3062, ".": 0.2524, "/": 0.3862, "0": 0.5068, "1": 0.5068,
    "2": 0.5068, "3": 0.5068, "4": 0.5068, "5": 0.5068, "6": 0.5068, "7": 0.5068,
    "8": 0.5068, "9": 0.5068, ":": 0.2676, ";": 0.2676, "<": 0.498, "=": 0.498,
    ">": 0.498, "?": 0.4634, "@": 0.894, "A": 0.5786, "B": 0.5439, "C": 0.5332,
    "D": 0.6152, "E": 0.4883, "F": 0.4595, "G": 0.6309, "H": 0.623, "I": 0.252,
    "J": 0.3188, "K": 0.5195, "L": 0.4204, "M": 0.855, "N": 0.6455, "O": 0.6621,
    "P": 0.5166, "Q": 0.6729, "R": 0.543, "S": 0.4595, "T": 0.4873, "U": 0.6416,
    "V": 0.5674, "W": 0.8896, "X": 0.519, "Y": 0.4873, "Z": 0.4683, "[": 0.3066,
    "\\": 0.3862, "]": 0.3066, "^": 0.498, "_": 0.498, "`": 0.291, "a": 0.479,
    "b": 0.5254, "c": 0.4229, "d": 0.5254, "e": 0.4976, "f": 0.3052, "g": 0.4707,
    "h": 0.5254, "i": 0.2295, "j": 0.2393, "k": 0.4546, "l": 0.2295, "m": 0.7988,
    "n": 0.5254, "o": 0.5273, "p": 0.5254, "q": 0.5254, "r": 0.3486, "s": 0.3911,
    "t": 0.335, "u": 0.5254, "v": 0.4517, "w": 0.7148, "x": 0.4331, "y": 0.4526,
    "z": 0.395, "{": 0.3145, "|": 0.4604, "}": 0.3145, "~": 0.498, "\u00a0": 0.2261,
    "\u00a9": 0.8345, "\u00b0": 0.3389, "\u00b1": 0.498, "\u00bc": 0.6362, "\u00bd": 0.6714, "\u00be": 0.6753,
    "\u00d7": 0.498, "\u00e9": 0.4976, "\u00f7": 0.498, "\u2013": 0.498, "\u2014": 0.9053, "\u2018": 0.2495,
    "\u2019": 0.2495, "\u201c": 0.4185, "\u201d": 0.4185, "\u2022": 0.498, "\u2026": 0.6904, "\u2192": 0.9053,
}
BOLD = {
    " ": 0.2261, "!": 0.3257, "\"": 0.4385, "#": 0.498, "$": 0.5068, "%": 0.729,
    "&": 0.7046, "'": 0.2334, "(": 0.3115, ")": 0.3115, "*": 0.498, "+": 0.498,
    ",": 0.2578, "-": 0.3062, ".": 0.2671, "/": 0.4297, "0": 0.5068, "1": 0.5068,
    "2": 0.5068, "3": 0.5068, "4": 0.5068, "5": 0.5068, "6": 0.5068, "7": 0.5068,
    "8": 0.5068, "9": 0.5068, ":": 0.2759, ";": 0.2759, "<": 0.498, "=": 0.498,
    ">": 0.498, "?": 0.4634, "@": 0.8984, "A": 0.606, "B": 0.5605, "C": 0.5293,
    "D": 0.6304, "E": 0.4878, "F": 0.459, "G": 0.6372, "H": 0.6309, "I": 0.2666,
    "J": 0.3311, "K": 0.5469, "L": 0.4229, "M": 0.874, "N": 0.6587, "O": 0.6763,
    "P": 0.5322, "Q": 0.686, "R": 0.563, "S": 0.4727, "T": 0.4951, "U": 0.6528,
    "V": 0.5913, "W": 0.9062, "X": 0.5508, "Y": 0.5195, "Z": 0.478, "[": 0.3247,
    "\\": 0.4297, "]": 0.3247, "^": 0.498, "_": 0.498, "`": 0.3003, "a": 0.4937,
    "b": 0.5366, "c": 0.4185, "d": 0.5366, "e": 0.5034, "f": 0.3164, "g": 0.4741,
    "h": 0.5366, "i": 0.2456, "j": 0.2554, "k": 0.48, "l": 0.2456, "m": 0.8135,
    "n": 0.5366, "o": 0.5376, "p": 0.5366, "q": 0.5366, "r": 0.3555, "s": 0.3989,
    "t": 0.3467, "u": 0.5366, "v": 0.4731, "w": 0.7451, "x": 0.4595, "y": 0.4736,
    "z": 0.3975, "{": 0.3438, "|": 0.4751, "}": 0.3438, "~": 0.498, "\u00a0": 0.2261,
    "\u00a9": 0.8345, "\u00b0": 0.3423, "\u00b1": 0.498, "\u00bc": 0.6577, "\u00bd": 0.6909, "\u00be": 0.7017,
    "\u00d7": 0.498, "\u00e9": 0.5034, "\u00f7": 0.498, "\u2013": 0.498, "\u2014": 0.9053, "\u2018": 0.2578,
    "\u2019": 0.2578, "\u201c": 0.4351, "\u201d": 0.4351, "\u2022": 0.498, "\u2026": 0.7114, "\u2192": 0.9053,
}
ITALIC = {
    " ": 0.2261, "!": 0.3257, "\"": 0.4009, "#": 0.498, "$": 0.5068, "%": 0.7148,
    "&": 0.6821, "'": 0.2207, "(": 0.3032, ")": 0.3032, "*": 0.498, "+": 0.498,
    ",": 0.2495, "-": 0.3062, ".": 0.2524, "/": 0.3877, "0": 0.5068, "1": 0.5068,
    "2": 0.5068, "3": 0.5068, "4": 0.5068, "5": 0.5068, "6": 0.5068, "7": 0.5068,
    "8": 0.5068, "9": 0.5068, ":": 0.2676, ";": 0.2676, "<": 0.498, "=": 0.498,
    ">": 0.498, "?": 0.4634, "@": 0.894, "A": 0.5786, "B": 0.5439, "C": 0.5225,
    "D": 0.6152, "E": 0.4883, "F": 0.4595, "G": 0.6309, "H": 0.623, "I": 0.252,
    "J": 0.3188, "K": 0.5195, "L": 0.4204, "M": 0.855, "N": 0.6445, "O": 0.6543,
    "P": 0.5166, "Q": 0.6641, "R": 0.543, "S": 0.4521, "T": 0.4873, "U": 0.6416,
    "V": 0.5674, "W": 0.8901, "X": 0.519, "Y": 0.4873, "Z": 0.4683, "[": 0.3066,
    "\\": 0.3843, "]": 0.3066, "^": 0.498, "_": 0.498, "`": 0.291, "a": 0.5142,
    "b": 0.5142, "c": 0.416, "d": 0.5142, "e": 0.4775, "f": 0.3052, "g": 0.5142,
    "h": 0.5142, "i": 0.2295, "j": 0.2393, "k": 0.4546, "l": 0.2295, "m": 0.791,
    "n": 0.5142, "o": 0.5132, "p": 0.5142, "q": 0.5142, "r": 0.3428, "s": 0.3892,
    "t": 0.335, "u": 0.5142, "v": 0.4458, "w": 0.7148, "x": 0.4331, "y": 0.4473,
    "z": 0.395, "{": 0.3145, "|": 0.4604, "}": 0.3145, "~": 0.498, "\u00a0": 0.2261,
    "\u00a9": 0.8345, "\u00b0": 0.3389, "\u00b1": 0.498, "\u00bc": 0.6362, "\u00bd": 0.6714, "\u00be": 0.6753,
    "\u00d7": 0.498, "\u00e9": 0.4775, "\u00f7": 0.498, "\u2013": 0.498, "\u2014": 0.9053, "\u2018": 0.2495,
    "\u2019": 0.2495, "\u201c": 0.4185, "\u201d": 0.4185, "\u2022": 0.498, "\u2026": 0.6899, "\u2192": 0.9053,
}
LINE = 1.2207      # (ascent + descent + line gap) / em, the same for all three
UNKNOWN = 0.60     # a character not in the tables: assume a wide one
EMU_PER_PT = 12700
SLACK_PT = 1.0     # rounding allowance before a box counts as over


def char_w(ch, size_pt, bold, italic):
    t = BOLD if bold else ITALIC if italic else REGULAR
    return t.get(ch, UNKNOWN) * size_pt


def paragraph_lines(p, width_pt, default_size):
    """Greedy word wrap of one paragraph. Returns (lines, line_height_pt, too_long_word)."""
    pieces = []  # (char, width) for the whole paragraph
    sizes = []
    for r in p.runs:
        size = r.font.size.pt if r.font.size else default_size
        sizes.append(size)
        for ch in r.text:
            pieces.append((ch, char_w(ch, size, bool(r.font.bold), bool(r.font.italic))))
    size = max(sizes) if sizes else default_size
    if not pieces:
        return 1, size * LINE, None
    words, cur = [], [[], 0.0]   # words with trailing space widths kept separate
    for ch, cw in pieces:
        if ch == " ":
            words.append((cur[1], cw, "".join(cur[0])))
            cur = [[], 0.0]
        else:
            cur[0].append(ch)
            cur[1] += cw
    words.append((cur[1], 0.0, "".join(cur[0])))
    lines, x, too_long = 1, 0.0, None
    for ww, sp, text in words:
        if ww > width_pt and too_long is None and text.strip(" "):
            too_long = text
        if x > 0 and x + ww > width_pt:
            lines += 1
            x = 0.0
        x += ww + sp
    return lines, size * LINE, too_long


def where(shape, sw, sh):
    """Name the quadrant a box sits in, Sherpa layout."""
    cx, cy = shape.left + shape.width / 2, shape.top + shape.height / 2
    if cy < sh * 0.12:
        return "header"
    if cy > sh * 0.96:
        return "footer"
    top, left = cy < sh * 0.53, cx < sw / 2
    return {(True, True): "quadrant iv", (True, False): "quadrant iii",
            (False, True): "quadrant i", (False, False): "quadrant ii"}[(top, left)]


def check(path):
    prs = Presentation(path)
    sw, sh = prs.slide_width, prs.slide_height
    problems = []
    for slide in prs.slides:
        for shape in slide.shapes:
            snippet = ""
            if shape.has_text_frame:
                snippet = " ".join(shape.text_frame.text.split())[:60]
            if (shape.left < 0 or shape.top < 0 or shape.left + shape.width > sw + 1
                    or shape.top + shape.height > sh + 1):
                problems.append({"where": where(shape, sw, sh), "text": snippet,
                                 "problem": "box runs off the page"})
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            tf = shape.text_frame
            ml, mr = tf.margin_left or 0, tf.margin_right or 0
            mt, mb = tf.margin_top or 0, tf.margin_bottom or 0
            avail_h = (shape.height - mt - mb) / EMU_PER_PT
            need_h = 0.0
            for p in tf.paragraphs:
                pPr = p._p.pPr
                indent = int(pPr.get("marL", 0)) if pPr is not None else 0
                width_pt = (shape.width - ml - mr - indent) / EMU_PER_PT
                n, lh, long_word = paragraph_lines(p, width_pt, 18)
                if long_word:
                    problems.append({"where": where(shape, sw, sh), "text": snippet,
                                     "problem": f"'{long_word[:30]}' is wider than its box"})
                ls = p.line_spacing
                if ls is not None and hasattr(ls, "pt"):
                    lh = ls.pt                      # exact spacing set by the builder
                elif isinstance(ls, float):
                    lh = lh * ls                    # a multiple, e.g. 0.9
                before = p.space_before.pt if p.space_before is not None else 0
                need_h += before + n * lh
            if need_h > avail_h + SLACK_PT:
                problems.append({"where": where(shape, sw, sh), "text": snippet,
                                 "problem": f"text needs {need_h:.0f}pt, box has {avail_h:.0f}pt",
                                 "over_pt": round(need_h - avail_h, 1)})
    return problems


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    probs = check(sys.argv[1])
    print(json.dumps({"ok": not probs, "problems": probs}, indent=2, ensure_ascii=False))
    sys.exit(0 if not probs else 3)
