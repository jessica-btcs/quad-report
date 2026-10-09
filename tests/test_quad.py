"""Run before every release: python -m unittest discover tests"""
import json, os, shutil, subprocess, sys, tempfile, unittest, warnings

warnings.simplefilter("ignore", ResourceWarning)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SKILL = os.path.join(ROOT, "plugins", "quad-report", "skills", "quad-report")
SAMPLE = os.path.join(SKILL, "examples", "sample-config.json")


def run(*args, ok=(0,)):
    p = subprocess.run([sys.executable, *args], capture_output=True, text=True, encoding="utf-8")
    if p.returncode not in ok:
        raise AssertionError(f"{args} exited {p.returncode}\n{p.stdout}\n{p.stderr}")
    return p


def data(*args, ok=(0,)):
    return json.loads(run(os.path.join(SKILL, "quad_data.py"), *args, ok=ok).stdout)


class Generator(unittest.TestCase):
    def test_sample_builds_one_slide_with_blank_quadrant_iv(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as d:
            out = os.path.join(d, "q.pptx")
            p = run(os.path.join(SKILL, "build_quad.py"), SAMPLE, out)
            self.assertIn("0 warning", p.stdout)
            prs = Presentation(out)
            self.assertEqual(len(prs.slides), 1)
            text = " ".join(s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame)
            for must in ("Quad Report 2026", "Revenue", "Sherpa Alliance", "What's On My Mind"):
                self.assertIn(must, text)

    def test_layout_check_passes_sample_and_refuses_overflow(self):
        cfg = json.load(open(SAMPLE, encoding="utf-8"))
        long = ("Measured the warehouse workflow with the operations manager team weekly " * 3)[:58]
        cfg["strategy_status"]["accomplishments"] = [long] * 3
        with tempfile.TemporaryDirectory() as d:
            c, out = os.path.join(d, "c.json"), os.path.join(d, "q.pptx")
            json.dump(cfg, open(c, "w", encoding="utf-8"))
            p = run(os.path.join(SKILL, "build_quad.py"), c, out, ok=(3,))
            self.assertIn("LAYOUT CHECK FAILED", p.stdout)
            self.assertIn("quadrant iii", p.stdout)
            self.assertIn("layout check: clean",
                          run(os.path.join(SKILL, "build_quad.py"), SAMPLE, out).stdout)

    def test_done_flag_renders(self):
        from pptx import Presentation
        cfg = json.load(open(SAMPLE, encoding="utf-8"))
        cfg["breakthrough_goals"][0]["flag"] = "done"
        with tempfile.TemporaryDirectory() as d:
            c, out = os.path.join(d, "c.json"), os.path.join(d, "q.pptx")
            json.dump(cfg, open(c, "w", encoding="utf-8"))
            run(os.path.join(SKILL, "build_quad.py"), c, out)
            text = " ".join(s.text_frame.text for s in Presentation(out).slides[0].shapes if s.has_text_frame)
            self.assertIn("DONE", text)


class DataFolder(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()
        self.dd = os.path.join(self.d, "Quad Report Data")

    def tearDown(self):
        shutil.rmtree(self.d)

    def cfg(self, period, revenue, goal_bullets=None):
        cfg = json.load(open(SAMPLE, encoding="utf-8"))
        cfg["period"] = period
        cfg["kpis"][0]["actual"] = revenue
        cfg["kpis"][0].pop("trend", None)
        for g in cfg["breakthrough_goals"]:
            g["flag"], g["stalled_months"] = None, None
        if goal_bullets:
            cfg["breakthrough_goals"][0]["bullets"] = goal_bullets
        path = os.path.join(self.d, f"{period}.json")
        json.dump(cfg, open(path, "w", encoding="utf-8"))
        return path

    def test_init_then_status_reports_first_run(self):
        data("init", self.dd)
        self.assertTrue(os.path.exists(os.path.join(self.dd, "README.md")))
        s = data("status", self.dd)
        self.assertFalse(s["profile_found"])
        self.assertEqual(s["months_on_file"], [])

    def test_missing_profile_fields_are_listed(self):
        data("init", self.dd)
        json.dump({"schema_version": 1, "company": {"name": "X"}},
                  open(os.path.join(self.dd, "quad-profile.json"), "w"))
        s = data("status", self.dd)
        self.assertTrue(s["needs_migration"])
        fields = [m["field"] for m in s["missing_fields"]]
        self.assertIn("company.report_by", fields)
        self.assertNotIn("company.name", fields)

    def test_archive_refuses_to_overwrite_and_keeps_old(self):
        data("init", self.dd)
        c = self.cfg("July 2026", 100000)
        self.assertTrue(data("archive", self.dd, c)["ok"])
        r = data("archive", self.dd, c, ok=(2,))
        self.assertFalse(r["ok"])
        r = data("archive", self.dd, c, "--replace")
        self.assertTrue(os.path.exists(r["previous_version_kept_as"]))

    def test_trends_and_stuck_goal_from_history(self):
        data("init", self.dd)
        same = ["Same words", "Same words again"]
        for period, rev in (("June 2026", 100000), ("July 2026", 110000), ("August 2026", 120000)):
            data("archive", self.dd, self.cfg(period, rev, same))
        sep = self.cfg("September 2026", 130000, same)
        r = data("annotate", self.dd, sep)
        self.assertEqual(r["period_key"], "2026-09")
        cfg = json.load(open(sep, encoding="utf-8"))
        self.assertEqual(cfg["kpis"][0]["trend"]["direction"], "up")
        self.assertEqual(cfg["kpis"][0]["trend"]["months"], 3)
        self.assertEqual(cfg["breakthrough_goals"][0]["flag"], "stalled")
        self.assertEqual(cfg["breakthrough_goals"][0]["stalled_months"], 3)
        self.assertEqual(data("status", self.dd)["suggested_period"], "2026-09")

    def test_reads_pre_release_history(self):
        data("init", self.dd)
        old = json.load(open(SAMPLE, encoding="utf-8"))
        old["period"] = "August 2026"
        json.dump(old, open(os.path.join(self.dd, "history", "2026-08.json"), "w"))
        self.assertEqual(data("status", self.dd)["last_month"], "2026-08")


class AnnualPlan(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()
        self.dd = os.path.join(self.d, "Quad Report Data")
        data("init", self.dd)

    def tearDown(self):
        shutil.rmtree(self.d)

    def status(self, prof, period):
        json.dump(prof, open(os.path.join(self.dd, "quad-profile.json"), "w"))
        return data("status", self.dd, "--period", period)

    def test_strategies_and_goals_are_required(self):
        fields = [m["field"] for m in self.status({"schema_version": 1}, "2026-09")["missing_fields"]]
        self.assertIn("strategies", fields)
        self.assertIn("goals", fields)

    def test_review_due_at_new_fiscal_year_only(self):
        prof = {"schema_version": 1, "plan_set": "2026-01", "company": {"fiscal_year_start_month": 1}}
        self.assertFalse(self.status(prof, "2026-09")["annual_review_due"])
        self.assertTrue(self.status(prof, "2027-01")["annual_review_due"])
        prof["company"]["fiscal_year_start_month"] = 7
        self.assertTrue(self.status(prof, "2026-07")["annual_review_due"])
        self.assertTrue(self.status({"schema_version": 1}, "2026-09")["annual_review_due"])


class IssueForm(unittest.TestCase):
    ISSUE = {"member": "Sample Owner", "meeting_date": "October 15, 2026", "group": "1234",
             "dollar_value": "$480K", "issue": "Crew leads.", "significant": "Margin.",
             "ideal_outcome": "Two leads.", "background": ["Started in March"],
             "options": ["Promote", "Hire"], "direction": "Promote.", "help": ["Options I'm missing?"]}

    def test_drawn_form_has_every_box_and_fills(self):
        import docx
        tool = os.path.join(SKILL, "fill_issue.py")
        with tempfile.TemporaryDirectory() as d:
            blank, issue, out = (os.path.join(d, n) for n in ("blank.docx", "issue.json", "out.docx"))
            run(tool, "blank", blank)
            self.assertTrue(json.loads(run(tool, "check", blank).stdout)["ok"])
            json.dump(self.ISSUE, open(issue, "w", encoding="utf-8"))
            r = json.loads(run(tool, "fill", issue, out).stdout)
            self.assertTrue(r["ok"])
            self.assertEqual(sorted(r["filled"]), sorted(self.ISSUE))
            text = " ".join(t.text or "" for t in docx.Document(out).element.body.iter(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
            for must in ("Crew leads.", "•  Promote", "OPPORTUNITY / ISSUE PROCESSING FORM"):
                self.assertIn(must, text)
            self.assertNotIn("Vistage", text)


if __name__ == "__main__":
    unittest.main()
